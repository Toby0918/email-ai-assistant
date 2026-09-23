"""Offline release smoke through the real window and local service."""
import json
import sqlite3
import sys
import tempfile
import time
import tkinter as tk
import urllib.request
from pathlib import Path

from .runtime import DesktopRuntime
from .window import DesktopWindow


def run_self_test(project: Path) -> dict:
    root = None
    runtime = None
    with tempfile.TemporaryDirectory(prefix="desktop-verification-", dir=project / "RuntimeTemp") as directory:
        try:
            runtime = DesktopRuntime(Path(directory), port=0)
            runtime.start()
            with urllib.request.urlopen(runtime.url + "/api/health") as response:
                assert json.load(response) == {"ok": True, "status": "ok"}
            for path in ("/", "/app.js", "/shared/render_analysis.js"):
                with urllib.request.urlopen(runtime.url + path) as response:
                    assert response.status == 200
                    assert len(response.read()) > 50
            root = tk.Tk()
            root.withdraw()
            window = DesktopWindow(root, runtime)
            window.load_sample()
            window.analyze()
            deadline = time.monotonic() + 15
            while window.busy and time.monotonic() < deadline:
                root.update()
                time.sleep(0.02)
            assert not window.busy
            assert window.completed_analysis is not None
            assert window.completed_analysis["reply_draft"]["needs_human_review"] is True
            assert window.draft.get("1.0", "end-1c").strip()
            assert "处理结论" in window.advice.get("1.0", "end-1c")
            window.copy_draft()
            assert "请先审核" in window.status.get()
            assert runtime.provider == "disabled"
            # The packaged application must parse a spreadsheet through its window,
            # not only import openpyxl successfully.
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active.append(["Quantity", "1200 pcs"])
            workbook.active.append(["Delivery", "Not confirmed; verify before replying"])
            spreadsheet_path = Path(directory) / "synthetic-order.xlsx"
            workbook.save(spreadsheet_path)
            workbook.close()
            window.paths = (spreadsheet_path,)
            window.set_text(window.body,
                            "Please confirm the delivery date. Please check with production "
                            "before promising any delivery date. No price has been agreed.", editable=True)
            window.analyze_button.invoke()
            deadline = time.monotonic() + 15
            while window.busy and time.monotonic() < deadline:
                root.update()
                time.sleep(0.02)
            assert not window.busy and window.completed_analysis is not None
            spreadsheet_insights = window.completed_analysis["attachment_insights"]
            assert len(spreadsheet_insights) == 1
            assert spreadsheet_insights[0]["status"] == "parsed"
            assert "Quantity: 1200 pcs" in spreadsheet_insights[0]["key_facts"]
            assert window.completed_analysis["category"] == "order_followup"
            assert not [fact for fact in window.completed_analysis["decision_brief"]["key_facts"]
                        if fact["label"] == "期限"]
            draft_body = window.draft.get("1.0", "end-1c")
            assert "1200 pcs" in draft_body and "before confirming any timing" in draft_body
            assert "Please confirm" not in draft_body and "Please check" not in draft_body
            root.update()
            window.reviewed.set(True)
            window.fields["subject"].set("Different synthetic email")
            window.copy_draft()
            assert window.completed_analysis is None and not window.reviewed.get()
            assert not window.draft.get("1.0", "end-1c").strip()
            assert "请先审核" in window.status.get()
            # State-notice regressions must survive freezing and UI presentation.
            window.paths = ()
            notice_samples = (
                ("Shipment status", "The shipment has been delivered to the receiving warehouse.", "已送达"),
                ("Payment update", "Our finance team is processing the payment. We will send the bank slip once available.", "付款处理中"),
                ("PPAP approval", "Please find attached the approved PPAP. The drawing has been updated to revision B.", "PPAP 已批准"),
            )
            for subject, body, expected in notice_samples:
                window.fields["subject"].set(subject)
                window.set_text(window.body, body, editable=True)
                window.analyze_button.invoke()
                deadline = time.monotonic() + 15
                while window.busy and time.monotonic() < deadline:
                    root.update()
                    time.sleep(0.02)
                assert not window.busy and window.completed_analysis is not None
                assert expected in window.advice.get("1.0", "end-1c")
                assert not window.completed_analysis["decision_brief"]["reply_recommendation"]["should_reply"]
                assert "Thank you for the update." in window.draft.get("1.0", "end-1c")
                assert not window.reviewed.get()
                window.copy_draft()
                assert "请先审核" in window.status.get()
            pricing_book = Workbook()
            pricing_book.active.append(["Material", "Crcy", "New Price", "Price Unit", "OPU", "Plant"])
            pricing_book.active.append(["SYN-731-CP", "USD", 8400, 1000, "EA", "SITE-A"])
            pricing_book.active.append(["SYN-732-BN", "INR", 624, 1, "KG", "SITE-B"])
            pricing_path = Path(directory) / "synthetic-pricing.xlsx"
            pricing_book.save(pricing_path)
            pricing_book.close()
            window.paths = (pricing_path,)
            window.fields["subject"].set("Synthetic pricing and inspection review")
            window.set_text(window.body, "Please review the attached table. Lot Qty.: 1400 Nos. "
                            "Salt spray passed. Dimensional inspection failed on two samples.", editable=True)
            window.analyze_button.invoke()
            deadline = time.monotonic() + 15
            while window.busy and time.monotonic() < deadline:
                root.update()
                time.sleep(0.02)
            assert not window.busy and window.completed_analysis is not None
            pricing_facts = window.completed_analysis['attachment_insights'][0]['key_facts']
            assert any('8.40 USD/EA' in fact and 'SYN-731-CP' in fact for fact in pricing_facts)
            assert any('624.00 INR/KG' in fact and 'SYN-732-BN' in fact for fact in pricing_facts)
            advice = window.advice.get('1.0', 'end-1c')
            assert '批次数量' in advice and '盐雾检验' in advice and '尺寸检验' in advice
            assert not window.reviewed.get()
            window.copy_draft()
            assert '请先审核' in window.status.get()
            from docx import Document
            document = Document()
            document.add_paragraph("Material: brass. Finish: chrome. MOQ: 1200 pcs.")
            document_path = Path(directory) / "synthetic-spec.docx"
            document.save(document_path)
            from .attachments import attachment_payload
            attachment_result = runtime.analyze({
                "subject": "Synthetic product specification", "from": "buyer@example.test",
                "body_text": "Please review the attached product specification.", "user_confirmed": True,
                "attachment_files": attachment_payload((document_path,), user_confirmed=True),
            })
            assert attachment_result["ok"]
            insights = attachment_result["analysis"]["attachment_insights"]
            assert len(insights) == 1 and insights[0]["status"] == "parsed"
            assert not list((Path(directory) / "attachment_temp").glob("*"))
            window.close()
            root = None
            runtime.close()
            runtime = DesktopRuntime(Path(directory), port=0)
            runtime.start()
            result = runtime.analyze({"subject": "Second synthetic request", "from": "buyer@example.test",
                                      "body_text": "Please confirm delivery date.", "user_confirmed": True})
            assert result["ok"] and result["saved_id"] >= 2
            # Libraries must survive freezing and upgrading without a provider call.
            from openai import OpenAI
            import openpyxl
            import docx
            import pypdf
            import PIL.Image
            client = OpenAI(api_key="synthetic-not-a-live-key", max_retries=0)
            client.close()
            return {"status": "PASS", "python": sys.version.split()[0], "sqlite": sqlite3.sqlite_version,
                    "native_window": "PASS", "sample_analysis": "PASS", "human_review_gate": "PASS",
                    "static_assets": "PASS", "restart_persistence": "PASS", "dependency_imports": "PASS",
                    "isolated_attachment_worker": "PASS", "attachment_cleanup": "PASS",
                    "xlsx_window_analysis": "PASS", "changed_email_copy_gate": "PASS",
                    "delivery_rule_semantics": "PASS",
                    "draft_request_echo_rejection": "PASS",
                    "business_notice_window_cases": 3,
                    "business_notice_window_semantics": "PASS",
                    "price_basis_and_scope_window_semantics": "PASS",
                    "provider_calls": 0, "live_mailbox_access": 0}
        finally:
            if root is not None:
                root.destroy()
            if runtime is not None:
                runtime.close()
