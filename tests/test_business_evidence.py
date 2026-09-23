"""Synthetic scope and table regressions through the submitted-email boundary."""
from datetime import UTC, datetime
from pathlib import Path
import tempfile
import unittest

from backend.email_agent.analyzer import analyze_current_email
from backend.email_agent.attachment_storage import StoredAttachment
from backend.email_agent.config import build_standalone_verification_config


class BusinessEvidenceTests(unittest.TestCase):
    def analyze_rows(self, rows):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "RuntimeTemp") as directory:
            path = Path(directory) / "synthetic-table.xlsx"
            book = Workbook()
            for row in rows:
                book.active.append(row)
            book.save(path)
            book.close()
            return self.analyze("Please review the attached table.", [
                StoredAttachment(path.name, 'xlsx', path, path.stat().st_size, datetime.now(UTC)),
            ])

    def analyze(self, body, attachments=()):
        return analyze_current_email({
            "subject": "Business review", "from": "buyer@example.test", "body_text": body,
            "stored_attachments": list(attachments),
        }, config=build_standalone_verification_config(
            sqlite_path=Path.cwd() / "RuntimeTemp/unused-evidence.sqlite3",
            attachment_temp_dir=Path.cwd() / "RuntimeTemp/unused-evidence-attachments",
        ))

    def test_price_per_thousand_retains_original_basis_and_computed_unit_price(self):
        result = self.analyze("Net price: USD 8400 per 1000 EA.")
        facts = str(result["decision_brief"]["key_facts"])
        self.assertIn("8400 USD per 1000 EA", facts)
        self.assertIn("8.40 USD/EA", facts)

    def test_selected_xlsx_binds_price_basis_and_item_on_each_row(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "RuntimeTemp") as directory:
            path = Path(directory) / "synthetic-pricing.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Material", "Currency", "Net Price", "Price Per Qty.", "Order Unit", "Purchase Org", "Plant"])
            sheet.append(["73101-CP", "USD", 8400, 1000, "EA", "ORG-A", "SITE-B"])
            sheet.append(["73102-BN", "EUR", 75, 10, "EA", "ORG-C", "SITE-D"])
            book.save(path)
            book.close()
            item = StoredAttachment(path.name, "xlsx", path, path.stat().st_size, datetime.now(UTC))
            result = self.analyze("Please review the attached pricing table.", [item])
            facts = result["attachment_insights"][0]["key_facts"]
            self.assertTrue(any("8.40 USD/EA" in f and "73101-CP" in f and "ORG-A" in f for f in facts), result["attachment_insights"])
            self.assertTrue(any("7.50 EUR/EA" in f and "73102-BN" in f and "ORG-C" in f for f in facts))
            self.assertFalse(any("8.40" in f and "73102-BN" in f for f in facts))

    def test_price_without_positive_basis_currency_and_unit_is_not_normalized(self):
        for text in ("Net price: USD 8400 per 0 EA.", "Net price: 8400 per 1000 EA.",
                     "Net price: USD 8400 per 1000.", "Net price: USD 8.400,00 per 1000 EA.",
                     "If approved, net price: USD 8400 per 1000 EA."):
            with self.subTest(text=text):
                self.assertNotIn("8.40 USD/EA", str(self.analyze(text)["decision_brief"]["key_facts"]))

    def test_quantity_roles_are_preserved_instead_of_one_undifferentiated_count(self):
        cases = (
            ("Lot Qty.: 1400 Nos. Please investigate the quality issue.", "批次数量", "1400", "不良数量"),
            ("Cargo gross weight: 3120 kg. VGM: 5340 kg.", "货物毛重", "3120 kg", "VGM"),
            ("Total inventory across all finishes is 52400 pcs. Please report the blocked quantity by finish.",
             "总库存", "52400 pcs", "各表面处理的冻结数量"),
        )
        for text, role, value, check in cases:
            with self.subTest(role=role):
                brief = self.analyze(text)["decision_brief"]
                self.assertTrue(any(item["label"] == role and value in item["value"] for item in brief["key_facts"]), brief)
                self.assertIn(check, str(brief))

    def test_allocation_proposal_retains_direction_and_does_not_authorize_execution(self):
        brief = self.analyze("We propose to take 240 units from PO 73101 and replenish them later from PO 73102. Kindly confirm.")["decision_brief"]
        facts = {item["label"]: item["value"] for item in brief["key_facts"]}
        self.assertIn("73101", facts.get("借出订单", ""))
        self.assertIn("73102", facts.get("补回订单", ""))
        self.assertIn("提议", str(brief["must_check"]))

    def test_unitless_lead_time_is_explicitly_unresolved(self):
        brief = self.analyze("Lead time: 36.")["decision_brief"]
        self.assertIn("交期单位", str(brief["missing_info"]))
        self.assertNotIn("36 days", str(brief))

    def test_different_inspection_projects_keep_separate_verdicts(self):
        brief = self.analyze("Salt spray passed. Dimensional inspection failed on two samples. Please measure six new samples.")["decision_brief"]
        facts = {item["label"]: item["value"] for item in brief["key_facts"]}
        self.assertIn("passed", facts.get("盐雾检验", ""))
        self.assertIn("failed", facts.get("尺寸检验", ""))
        self.assertIn("two samples", facts.get("尺寸检验", ""))
        self.assertIn("样本", str(brief["must_check"]))

    def test_conditional_result_is_not_unconditional_approval_or_a_template_option(self):
        result = self.analyze("Please see the updated ISIR. The result is Conditional OK; approval is subject to the stated conditions.")
        self.assertIn("条件接受", result["decision_brief"]["one_line_conclusion"])
        self.assertIn("条件接受", result["summary"])
        self.assertNotIn("complete the internal review", result["reply_draft"]["body"])
        self.assertIn("stated conditions", result["reply_draft"]["body"])
        for text in ("Options: Conditional OK / NG / OK.", "The result is not Conditional OK.",
                     "If the result is Conditional OK, please notify us.",
                     "Not confirmed:\nThe result is Conditional OK.",
                     "The result is Conditional OK. The result is NG."):
            with self.subTest(text=text):
                brief = self.analyze(text)["decision_brief"]
                self.assertFalse(any(item["label"] == "条件检验结论" for item in brief["key_facts"]))

    def test_actual_header_variants_and_mixed_bases_survive_the_selected_file_parser(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "RuntimeTemp") as directory:
            path = Path(directory) / "synthetic-pir.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Material Number", "P. Org", "Plant", "New Price", "Crcy", "Price Per Qty.", "OPU", "Effective Date"])
            sheet.append(["SYN-711-CP", "ORG-A", "PLANT-X", 7650, "USD", 1000, "EA", "2026-10-01"])
            sheet.append(["Material", "Purch Org", "Plant", "Q4 2026", "IR Cur", "Price Unit", "Ord Un"])
            sheet.append(["SYN-712", "ORG-A", "PLANT-Y", 15.7, "USD", 1000, "EA"])
            sheet.append(["SYN-712", "ORG-B", "PLANT-Z", 1.26, "USD", 1, "EA"])
            book.save(path)
            book.close()
            result = self.analyze("Please review the attached table.", [StoredAttachment(path.name, 'xlsx', path, path.stat().st_size, datetime.now(UTC))])
            facts = result['attachment_insights'][0]['key_facts']
            self.assertTrue(any('7.65 USD/EA' in f and '2026-10-01' in f for f in facts), facts)
            self.assertTrue(any('0.015700 USD/EA' in f and 'PLANT-Y' in f for f in facts), facts)
            self.assertTrue(any('1.26 USD/EA' in f and 'PLANT-Z' in f for f in facts), facts)
            self.assertFalse(any(f.startswith('Amount:') for f in facts), facts)

    def test_incomplete_new_header_cannot_inherit_previous_price_columns(self):
        result = self.analyze_rows([
            ["Material", "Currency", "Net Price", "Price Unit", "Order Unit"],
            ["SKU-711", "USD", 8400, 1000, "EA"],
            ["Description", "Curr.", "Gross Price", "Price Unit", "UOM"],
            ["SKU-712", "USD", 10000, 1000, "EA"],
        ])
        facts = result['attachment_insights'][0]['key_facts']
        self.assertTrue(any('8.40 USD/EA' in f for f in facts))
        self.assertFalse(any('10.00 USD/EA' in f for f in facts), facts)

    def test_retracted_price_and_template_status_are_not_promoted(self):
        for text in (
            "Net price: USD 8400 per 1000 EA. Correction: this price was incorrect; no price is agreed.",
            "Options:\nThe result is Conditional OK.\nPlease select one of these template options.",
        ):
            with self.subTest(text=text):
                facts = str(self.analyze(text)['decision_brief']['key_facts'])
                self.assertNotIn('8.40 USD/EA', facts)
                self.assertNotIn('条件检验结论', facts)

    def test_approval_and_pending_finish_have_distinct_facts_and_consistent_draft(self):
        result = self.analyze("PPAP for SKU 73131-BN is approved. PPAP for SKU 73131-CP is pending.")
        facts = {item['label']: item['value'] for item in result['decision_brief']['key_facts']}
        self.assertEqual(facts.get('PPAP 批准料号'), '73131-BN')
        self.assertEqual(facts.get('PPAP 待批准料号'), '73131-CP')
        self.assertIn('范围', result['summary'])
        self.assertNotIn('complete the internal review', result['reply_draft']['body'])

    def test_scoped_status_preserves_quality_request_and_rca_draft(self):
        for status in ('PPAP for SKU 73131-BN is approved.', 'The result is Conditional OK.'):
            with self.subTest(status=status):
                result = self.analyze(status + ' Please investigate the defective products and provide an RCA by tomorrow.')
                self.assertIn('质量', result['summary'])
                self.assertIn('RCA', result['reply_draft']['body'])

    def test_cancelled_price_is_not_a_current_price_fact(self):
        result = self.analyze('Net price: USD 8400 per 1000 EA. This price has been cancelled. Do not use it.')
        self.assertNotIn('8.40 USD/EA', str(result['decision_brief']['key_facts']))

    def test_price_scope_rejects_sensitive_identifier_shapes(self):
        for token in ('202-555-0199', '13800138000', 'portal.example.test'):
            with self.subTest(token=token):
                result = self.analyze_rows([
                    ['Material', 'Currency', 'Net Price', 'Price Unit', 'Order Unit'],
                    [token, 'USD', 8400, 1000, 'EA'],
                ])
                self.assertNotIn(token, str(result['attachment_insights'][0]['key_facts']))

    def test_raw_constructed_price_line_does_not_impersonate_parser_output(self):
        for suffix in ('This price was withdrawn.', 'Unverified source example.'):
            result = self.analyze_rows([['Price basis: 100 USD per 10 EA = 10.00 USD/EA'], [suffix]])
            self.assertFalse(any(f.startswith('Price basis:') for f in result['attachment_insights'][0]['key_facts']))

    def test_short_alphabetic_plant_keeps_row_scope(self):
        result = self.analyze_rows([
            ['Material', 'Currency', 'Net Price', 'Price Unit', 'Order Unit', 'Plant'],
            ['SYN-711', 'USD', 8400, 1000, 'EA', 'CE'],
        ])
        self.assertTrue(any('8.40 USD/EA' in f and 'Plant: CE' in f
                            for f in result['attachment_insights'][0]['key_facts']))


if __name__ == "__main__":
    unittest.main()
